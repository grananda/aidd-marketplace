#!/usr/bin/env python3
"""Genera el deck ejecutivo de evolución de Native AI y su PDF de revisión."""

from __future__ import annotations

import json
import math
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent
RENDERED = OUT / "_rendered"
PPTX_PATH = OUT / "evolucion-native-ai-v2.pptx"
PDF_PATH = OUT / "evolucion-native-ai-v2.pdf"
PREVIEW_PATH = OUT / "preview.png"

W_IN, H_IN = 13.333, 7.5
W_PX, H_PX = 1600, 900
SX, SY = W_PX / W_IN, H_PX / H_IN

FONT_REGULAR = "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/noto/NotoSans-Bold.ttf"
FONT_LIGHT = "/usr/share/fonts/truetype/noto/NotoSans-Light.ttf"
PPT_FONT = "Arial"

C = {
    "paper": "F6F4EF",
    "white": "FFFFFF",
    "ink": "17324D",
    "ink2": "3F556A",
    "muted": "71808E",
    "line": "D7DEE3",
    "blue": "0072BC",
    "blue2": "DCEEF8",
    "cyan": "16A9C7",
    "teal": "008F83",
    "teal2": "DDF2EE",
    "amber": "E8792E",
    "amber2": "FCE9DC",
    "yellow": "F4B942",
    "yellow2": "FFF4D2",
    "navy": "0D2942",
    "soft": "E9ECEF",
    "grey": "A9B4BE",
    "red": "C94B48",
    "red2": "F8E3E2",
}


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def rgba(value: str, alpha: int = 255) -> tuple[int, int, int, int]:
    value = value.lstrip("#")
    return int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16), alpha


def font_path(bold: bool = False, light: bool = False) -> str:
    candidate = FONT_BOLD if bold else FONT_LIGHT if light else FONT_REGULAR
    if Path(candidate).exists():
        return candidate
    return "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"


def emu(v: float):
    return Inches(v)


@dataclass
class TextStyle:
    size: float
    color: str = C["ink"]
    bold: bool = False
    light: bool = False
    align: str = "left"
    valign: str = "top"
    min_size: float = 8
    line_spacing: float = 1.12


class Canvas:
    def __init__(self, prs: Presentation, number: int, section: str, title: str | None = None, subtitle: str | None = None, dark: bool = False):
        self.prs = prs
        self.slide = prs.slides.add_slide(prs.slide_layouts[6])
        self.number = number
        self.dark = dark
        self.image = Image.new("RGB", (W_PX, H_PX), rgba(C["navy"] if dark else C["paper"])[0:3])
        self.draw = ImageDraw.Draw(self.image)
        self.rect(0, 0, W_IN, H_IN, C["navy"] if dark else C["paper"], radius=0)
        if title:
            self.header(section, title, subtitle)

    @staticmethod
    def _px(x: float, y: float) -> tuple[int, int]:
        return round(x * SX), round(y * SY)

    def rect(self, x: float, y: float, w: float, h: float, fill: str, line: str | None = None, radius: float = 0.08, line_width: float = 1):
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shape = self.slide.shapes.add_shape(shape_type, emu(x), emu(y), emu(w), emu(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape.line.color.rgb = rgb(line or fill)
        shape.line.width = Pt(line_width)
        x0, y0 = self._px(x, y)
        x1, y1 = self._px(x + w, y + h)
        if radius:
            self.draw.rounded_rectangle((x0, y0, x1, y1), radius=round(radius * SX), fill=rgba(fill), outline=rgba(line or fill), width=max(1, round(line_width * 1.3)))
        else:
            self.draw.rectangle((x0, y0, x1, y1), fill=rgba(fill), outline=rgba(line or fill), width=max(1, round(line_width * 1.3)))
        return shape

    def line(self, x1: float, y1: float, x2: float, y2: float, color: str, width: float = 1, dash: bool = False):
        shape = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(x1), emu(y1), emu(max(0.002, x2 - x1)), emu(max(0.002, y2 - y1)))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(color)
        shape.line.fill.background()
        px1, py1 = self._px(x1, y1)
        px2, py2 = self._px(x2, y2)
        self.draw.line((px1, py1, px2, py2), fill=rgba(color), width=max(1, round(width * 1.6)))

    def circle(self, x: float, y: float, d: float, fill: str, line: str | None = None):
        shape = self.slide.shapes.add_shape(MSO_SHAPE.OVAL, emu(x), emu(y), emu(d), emu(d))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape.line.color.rgb = rgb(line or fill)
        x0, y0 = self._px(x, y)
        x1, y1 = self._px(x + d, y + d)
        self.draw.ellipse((x0, y0, x1, y1), fill=rgba(fill), outline=rgba(line or fill))

    def arrow(self, x: float, y: float, w: float, h: float, fill: str):
        shape = self.slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, emu(x), emu(y), emu(w), emu(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape.line.fill.background()
        x0, y0 = self._px(x, y)
        x1, y1 = self._px(x + w, y + h)
        head = round(w * SX * 0.25)
        mid = (y0 + y1) // 2
        pts = [(x0, y0 + (y1 - y0) // 4), (x1 - head, y0 + (y1 - y0) // 4), (x1 - head, y0), (x1, mid), (x1 - head, y1), (x1 - head, y1 - (y1 - y0) // 4), (x0, y1 - (y1 - y0) // 4)]
        self.draw.polygon(pts, fill=rgba(fill))

    def pill(self, x: float, y: float, w: float, h: float, text: str, fill: str, color: str, size: float = 10, bold: bool = True):
        self.rect(x, y, w, h, fill, radius=h / 2)
        self.text(x + 0.08, y + 0.02, w - 0.16, h - 0.04, text, TextStyle(size, color, bold, align="center", valign="middle", min_size=7))

    def _wrap(self, value: str, pil_font: ImageFont.FreeTypeFont, max_width: int) -> str:
        lines: list[str] = []
        for paragraph in value.split("\n"):
            if not paragraph:
                lines.append("")
                continue
            words = paragraph.split(" ")
            current = words[0]
            for word in words[1:]:
                trial = current + " " + word
                if self.draw.textlength(trial, font=pil_font) <= max_width:
                    current = trial
                else:
                    lines.append(current)
                    current = word
            lines.append(current)
        return "\n".join(lines)

    def _fit(self, value: str, w: float, h: float, style: TextStyle) -> tuple[float, str, ImageFont.FreeTypeFont]:
        max_w, max_h = round(w * SX), round(h * SY)
        size = style.size
        while size >= style.min_size:
            pil_font = ImageFont.truetype(font_path(style.bold, style.light), round(size * 1.62))
            wrapped = self._wrap(value, pil_font, max_w)
            box = self.draw.multiline_textbbox((0, 0), wrapped, font=pil_font, spacing=round(size * style.line_spacing * 0.46), align=style.align)
            if box[2] - box[0] <= max_w and box[3] - box[1] <= max_h:
                return size, wrapped, pil_font
            size -= 0.5
        pil_font = ImageFont.truetype(font_path(style.bold, style.light), round(style.min_size * 1.62))
        return style.min_size, self._wrap(value, pil_font, max_w), pil_font

    def text(self, x: float, y: float, w: float, h: float, value: str, style: TextStyle):
        chosen, wrapped, pil_font = self._fit(value, w, h, style)
        tb = self.slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
        tf = tb.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[style.valign]
        lines = value.split("\n")
        for idx, line in enumerate(lines):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = PPT_FONT
            p.font.size = Pt(chosen)
            p.font.bold = style.bold
            p.font.color.rgb = rgb(style.color)
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[style.align]
            p.space_after = Pt(chosen * 0.12)
        px, py = self._px(x, y)
        max_w, max_h = round(w * SX), round(h * SY)
        box = self.draw.multiline_textbbox((0, 0), wrapped, font=pil_font, spacing=round(chosen * style.line_spacing * 0.46), align=style.align)
        tw, th = box[2] - box[0], box[3] - box[1]
        tx = px if style.align == "left" else px + (max_w - tw) // 2 if style.align == "center" else px + max_w - tw
        ty = py if style.valign == "top" else py + (max_h - th) // 2 if style.valign == "middle" else py + max_h - th
        self.draw.multiline_text((tx, ty), wrapped, font=pil_font, fill=rgba(style.color), spacing=round(chosen * style.line_spacing * 0.46), align=style.align)
        return chosen

    def header(self, section: str, title: str, subtitle: str | None):
        self.text(0.55, 0.3, 3.2, 0.25, section.upper(), TextStyle(9, C["blue"], True))
        self.text(0.55, 0.72, 12.1, 0.58, title, TextStyle(25, C["ink"], True, min_size=20))
        if subtitle:
            self.text(0.55, 1.28, 12.0, 0.45, subtitle, TextStyle(10.5, C["ink2"], min_size=9))
        self.line(0.55, 1.78, 12.78, 1.78, C["line"], 1)

    def footer(self, total: int, source: str = "Fuente: repositorio aidd-marketplace · VERSION 1.46.1"):
        color = C["grey"] if self.dark else C["muted"]
        self.text(0.55, 7.14, 9.8, 0.18, source, TextStyle(6.8, color, min_size=6))
        self.text(11.75, 7.10, 1.0, 0.22, f"{self.number:02d} / {total:02d}", TextStyle(7.2, color, True, align="right", min_size=6))


def metrics() -> dict:
    marketplace = json.loads((ROOT / ".claude-plugin/marketplace.json").read_text(encoding="utf-8"))
    plugins = marketplace["plugins"]
    skill_counts = {}
    for plugin in plugins:
        plugin_root = (ROOT / plugin["source"]).resolve()
        skill_counts[plugin["name"]] = len(list(plugin_root.glob("skills/*/SKILL.md")))
    command_counts = {"aidd": 9, "aisdd": 9, "aiba": 7, "aiad": 11, "aifg": 2, "boosters": 3}
    result = {
        "plugins": len(plugins),
        "skills": sum(skill_counts.values()),
        "commands": sum(command_counts.values()),
        "checks": len(list((ROOT / ".github/scripts").glob("check_*.py"))),
        "topologies": 3,
        "skill_counts": skill_counts,
        "command_counts": command_counts,
    }
    expected = {"plugins": 6, "skills": 34, "commands": 41, "checks": 12, "topologies": 3}
    for key, value in expected.items():
        if result[key] != value:
            raise RuntimeError(f"Métrica {key} cambió: esperada {value}, observada {result[key]}. Revisa el relato antes de regenerar.")
    return result


def card(c: Canvas, x: float, y: float, w: float, h: float, eyebrow: str, title: str, body: str, accent: str = C["blue"], fill: str = C["white"]):
    c.rect(x, y, w, h, fill, C["line"], radius=0.12)
    c.rect(x, y, 0.07, h, accent, accent, radius=0)
    c.text(x + 0.24, y + 0.18, w - 0.42, 0.22, eyebrow.upper(), TextStyle(8.5, accent, True))
    c.text(x + 0.24, y + 0.48, w - 0.42, 0.50, title, TextStyle(15, C["ink"], True, min_size=11))
    c.text(x + 0.24, y + 1.03, w - 0.42, h - 1.18, body, TextStyle(9.2, C["ink2"], min_size=7.5))


def metric_card(c: Canvas, x: float, y: float, w: float, label: str, before: str, after: str, accent: str):
    c.rect(x, y, w, 1.36, C["white"], C["line"], radius=0.12)
    c.text(x + 0.18, y + 0.15, w - 0.36, 0.25, label.upper(), TextStyle(7.8, C["muted"], True, align="center"))
    c.text(x + 0.22, y + 0.46, 0.55, 0.55, before, TextStyle(19, C["grey"], True, align="center", valign="middle"))
    c.arrow(x + 0.82, y + 0.61, 0.45, 0.19, accent)
    c.text(x + 1.31, y + 0.42, w - 1.48, 0.62, after, TextStyle(23, accent, True, align="center", valign="middle"))
    c.text(x + 0.18, y + 1.09, w - 0.36, 0.16, "PARTIDA  →  HOY", TextStyle(6.4, C["muted"], True, align="center"))


def build_slides(prs: Presentation, m: dict) -> list[Canvas]:
    slides: list[Canvas] = []
    total = 18

    # 1 — Cover
    c = Canvas(prs, 1, "", dark=True)
    c.rect(0, 0, 0.16, H_IN, C["cyan"], radius=0)
    c.text(0.7, 0.52, 4.8, 0.25, "NTT DATA SPAIN · GDN-e", TextStyle(9, C["cyan"], True))
    c.text(0.7, 1.20, 8.9, 1.55, "Native AI ya no es solo un método para construir", TextStyle(34, C["white"], True, min_size=29))
    c.text(0.7, 2.90, 8.8, 0.82, "Es un sistema operativo de entrega: conecta lo que pide el cliente, lo que construye el equipo y la evidencia con la que se decide.", TextStyle(15, "C6D6E3", min_size=12))
    c.rect(9.85, 0.72, 2.72, 5.72, "143B59", "28516F", radius=0.18)
    stages = [("01", "ESPECIFICAR", C["blue2"], C["blue"]), ("02", "ENTREGAR", C["teal2"], C["teal"]), ("03", "MEDIR", C["amber2"], C["amber"])]
    for i, (n, label, fill, color) in enumerate(stages):
        yy = 1.18 + i * 1.58
        c.circle(10.23, yy, 0.56, fill)
        c.text(10.23, yy + 0.02, 0.56, 0.48, n, TextStyle(11, color, True, align="center", valign="middle"))
        c.text(10.98, yy + 0.07, 1.25, 0.32, label, TextStyle(11, C["white"], True, valign="middle"))
        if i < 2:
            c.line(10.51, yy + 0.62, 10.51, yy + 1.42, "4B6A82", 2)
    c.pill(0.7, 5.50, 2.08, 0.42, "RESUMEN EJECUTIVO", C["blue"], C["white"], 8.2)
    c.text(0.7, 6.10, 7.8, 0.42, "Evolución de Native AI · metodología original → marketplace v1.46.1", TextStyle(10.5, "AFC2D1"))
    c.footer(total, "Comparativa verificada sobre metodología v2.0, native-ai-specs v1.6.0 y marketplace v1.46.1")
    slides.append(c)

    # 2 — thesis
    c = Canvas(prs, 2, "Resumen ejecutivo", "La evolución cabe en una frase", "El núcleo no cambió. Cambió la distancia que recorre.")
    c.rect(0.62, 2.15, 3.42, 3.75, C["white"], C["line"], radius=0.16)
    c.text(0.90, 2.45, 2.85, 0.30, "PUNTO DE PARTIDA", TextStyle(9, C["blue"], True))
    c.text(0.90, 2.93, 2.85, 0.96, "Especificar antes de implementar", TextStyle(23, C["ink"], True, min_size=18))
    c.text(0.90, 4.10, 2.78, 1.25, "OpenSpec, pre-flight humano, changes pequeños y auditoría: el método resolvía la deriva entre requisitos y código.", TextStyle(11, C["ink2"], min_size=9))
    c.arrow(4.27, 3.48, 1.18, 0.52, C["amber"])
    c.text(4.18, 4.15, 1.35, 0.48, "EXTENDER\nSIN ROMPER", TextStyle(8, C["amber"], True, align="center"))
    c.rect(5.70, 2.15, 6.98, 3.75, C["navy"], C["navy"], radius=0.16)
    c.text(6.05, 2.45, 3.0, 0.30, "RESULTADO", TextStyle(9, C["cyan"], True))
    c.text(6.05, 2.93, 5.95, 0.92, "Una cadena completa de entrega", TextStyle(25, C["white"], True, min_size=20))
    chain = [("CLIENTE", C["blue"]), ("DISEÑO", C["cyan"]), ("CÓDIGO", C["teal"]), ("EVIDENCIA", C["amber"])]
    for i, (label, color) in enumerate(chain):
        xx = 6.05 + i * 1.52
        c.pill(xx, 4.22, 1.20, 0.42, label, color, C["white"], 7.5)
        if i < len(chain) - 1:
            c.arrow(xx + 1.23, 4.34, 0.24, 0.16, "7292A9")
    c.text(6.05, 4.92, 5.90, 0.62, "La misma fuente de verdad llega hasta el entregable que se firma y el KPI que se defiende.", TextStyle(11, "C6D6E3", min_size=9))
    c.rect(0.62, 6.20, 12.06, 0.56, C["amber2"], C["amber2"], radius=0.08)
    c.text(0.90, 6.31, 11.50, 0.29, "TESIS  ·  No es otro método con el mismo nombre: es el mismo principio llevado a los extremos donde un proyecto real lo necesita.", TextStyle(10.5, C["ink"], True, align="center"))
    c.footer(total)
    slides.append(c)

    # 3 — metrics
    c = Canvas(prs, 3, "Resumen ejecutivo", "El salto, medido sobre los artefactos", "Cinco indicadores de superficie, cobertura e industrialización. Ninguno es una estimación.")
    metric_card(c, 0.62, 2.15, 2.26, "Skills", "3", str(m["skills"]), C["blue"])
    metric_card(c, 3.03, 2.15, 2.26, "Comandos", "7", str(m["commands"]), C["cyan"])
    metric_card(c, 5.44, 2.15, 2.26, "Fases con herramienta", "3/8", "8/8", C["teal"])
    metric_card(c, 7.85, 2.15, 2.26, "Topologías", "1", str(m["topologies"]), C["amber"])
    metric_card(c, 10.26, 2.15, 2.26, "Checks CI", "0", str(m["checks"]), C["red"])
    c.text(0.62, 3.85, 4.20, 0.34, "QUÉ DICEN REALMENTE LAS CIFRAS", TextStyle(9, C["blue"], True))
    statements = [
        ("Cobertura", "La automatización ya acompaña al proyecto completo, no solo a la ejecución."),
        ("Modularidad", "La capacidad se instala por necesidad: seis plugins, no una suite monolítica."),
        ("Control", "El método se valida en CI y deja evidencia auditable de cada transición."),
    ]
    for i, (t, b) in enumerate(statements):
        xx = 0.62 + i * 4.05
        c.rect(xx, 4.35, 3.82, 1.55, C["white"], C["line"], radius=0.10)
        c.circle(xx + 0.20, 4.62, 0.36, [C["blue2"], C["teal2"], C["amber2"]][i])
        c.text(xx + 0.20, 4.62, 0.36, 0.34, str(i + 1), TextStyle(9, [C["blue"], C["teal"], C["amber"]][i], True, align="center", valign="middle"))
        c.text(xx + 0.72, 4.52, 2.82, 0.30, t, TextStyle(12, C["ink"], True))
        c.text(xx + 0.72, 4.88, 2.82, 0.75, b, TextStyle(8.8, C["ink2"], min_size=7.5))
    c.footer(total, "Conteo: SKILL.md, comandos publicados, topologías AISDD y check_*.py · 09/09/2026")
    slides.append(c)

    # 4 — lifecycle coverage
    c = Canvas(prs, 4, "Cobertura", "La herramienta ya no empieza a mitad del proyecto", "La metodología describía el recorrido; el marketplace convierte cada tramo en una capacidad ejecutable.")
    stages = ["Brief", "Requisitos", "Diseño", "Plan", "Build", "Validación", "Entrega", "Medición"]
    colors = [C["blue"], C["blue"], C["cyan"], C["cyan"], C["teal"], C["teal"], C["amber"], C["amber"]]
    x0, y0, cellw = 2.72, 2.33, 1.18
    for i, label in enumerate(stages):
        c.text(x0 + i * cellw, 2.00, cellw - 0.06, 0.28, label.upper(), TextStyle(7.2, C["muted"], True, align="center"))
    rows = [
        ("Método original", "definía", [1, 1, 1, 1, 1, 1, 1, 0], C["grey"]),
        ("Tooling v1.6", "automatizaba", [0, 0, 0, 1, 1, 1, 0, 0], C["blue"]),
        ("Marketplace", "ejecuta", [1, 1, 1, 1, 1, 1, 1, 1], C["teal"]),
    ]
    for r, (name, verb, coverage, accent) in enumerate(rows):
        yy = y0 + r * 1.0
        c.text(0.65, yy + 0.05, 1.95, 0.30, name, TextStyle(12, C["ink"], True))
        c.text(0.65, yy + 0.38, 1.95, 0.22, verb, TextStyle(8, C["muted"]))
        for i, active in enumerate(coverage):
            fill = colors[i] if active and r == 2 else accent if active else C["soft"]
            c.rect(x0 + i * cellw, yy, cellw - 0.09, 0.68, fill, fill, radius=0.10)
            if active:
                c.text(x0 + i * cellw, yy + 0.15, cellw - 0.09, 0.30, "SÍ", TextStyle(8.5, C["white"], True, align="center", valign="middle"))
    c.rect(0.65, 5.52, 12.0, 0.95, C["white"], C["line"], radius=0.12)
    c.text(0.93, 5.74, 2.28, 0.40, "La diferencia clave", TextStyle(12, C["amber"], True, valign="middle"))
    c.text(3.15, 5.70, 9.05, 0.50, "No se inventaron requisitos, historias o arquitectura: se transformaron guiones manuales en contratos ejecutables y salidas consistentes.", TextStyle(11, C["ink"], min_size=9, valign="middle"))
    c.footer(total, "Criterio de fase: 8 tramos operativos de brief a medición; ‘definía’ no equivale a ‘tenía herramienta’.")
    slides.append(c)

    # 5 — layers
    c = Canvas(prs, 5, "Arquitectura de la evolución", "Tres capas; una sola línea de continuidad", "Cada capa conserva la anterior y amplía la promesa que puede cumplir.")
    layers = [
        (0.72, 4.98, 11.90, 1.15, C["navy"], C["white"], "1 · METODOLOGÍA", "Especificación como motor · documentos como verdad · human-in-the-loop · roles y handoffs"),
        (1.48, 3.62, 10.38, 1.15, C["blue"], C["white"], "2 · PROCEDIMIENTO EJECUTABLE", "OpenSpec · pre-flight · roadmap consciente del contexto · auditoría obligatoria"),
        (2.24, 2.26, 8.86, 1.15, C["teal"], C["white"], "3 · SISTEMA DE ENTREGA", "Cliente y negocio · dos motores de ejecución · diseño real · medición · CI · multirepo"),
    ]
    for x, y, w, h, fill, color, t, b in layers:
        c.rect(x, y, w, h, fill, fill, radius=0.14)
        c.text(x + 0.30, y + 0.18, 3.15, 0.28, t, TextStyle(9, color, True))
        c.text(x + 3.25, y + 0.16, w - 3.55, 0.58, b, TextStyle(10.2, color, min_size=8.5, valign="middle"))
    c.pill(9.60, 2.57, 1.18, 0.36, "NUEVO", C["amber2"], C["amber"], 7.5)
    c.pill(10.34, 3.93, 1.18, 0.36, "EJECUTA", C["blue2"], C["blue"], 7.5)
    c.pill(11.08, 5.29, 1.18, 0.36, "CONSERVA", "31506A", C["white"], 7.5)
    c.text(0.72, 6.43, 11.80, 0.35, "Lectura ejecutiva: el marketplace no compite con Native AI; convierte su filosofía en una capacidad operativa, distribuible y gobernable.", TextStyle(11, C["ink"], True, align="center"))
    c.footer(total)
    slides.append(c)

    # 6 — ecosystem
    c = Canvas(prs, 6, "Ecosistema", "Seis plugins, organizados alrededor de una única fuente de verdad", "La modularidad permite instalar solo la superficie que el proyecto necesita.")
    c.rect(4.45, 2.58, 4.42, 2.15, C["navy"], C["navy"], radius=0.18)
    c.text(4.82, 2.90, 3.68, 0.30, "FUENTE DE VERDAD", TextStyle(9, C["cyan"], True, align="center"))
    c.text(4.82, 3.35, 3.68, 0.60, "HU · arquitectura · specs · decisiones", TextStyle(18, C["white"], True, align="center", valign="middle", min_size=14))
    c.text(4.82, 4.10, 3.68, 0.25, "Documentos versionados", TextStyle(9, "B8CBD9", align="center"))
    plugin_cards = [
        (0.62, 2.05, 3.22, "AIDD", "Define el qué", "brief → requisitos → diseño", C["blue"]),
        (0.62, 4.55, 3.22, "AIBA", "Da la cara al negocio", "firma → plan → estado → KPI", C["amber"]),
        (9.48, 2.05, 3.22, "AISDD", "Ejecuta con IA", "change → código → validación", C["teal"]),
        (9.48, 4.55, 3.22, "AIAD", "Aumenta al humano", "pensar → escribir → aprender", C["cyan"]),
        (4.08, 5.32, 2.38, "AIFG", "Conecta diseño real", "Figma → componente → HU", C["red"]),
        (6.86, 5.32, 2.38, "BOOSTERS", "Visualiza", "UX · UML · HTML", C["ink2"]),
    ]
    for x, y, w, name, purpose, flow, accent in plugin_cards:
        c.rect(x, y, w, 1.35, C["white"], C["line"], radius=0.12)
        c.rect(x, y, 0.08, 1.35, accent, accent, radius=0)
        c.text(x + 0.26, y + 0.17, w - 0.44, 0.25, name, TextStyle(9, accent, True))
        c.text(x + 0.26, y + 0.49, w - 0.44, 0.30, purpose, TextStyle(12.5, C["ink"], True, min_size=10))
        c.text(x + 0.26, y + 0.92, w - 0.44, 0.22, flow, TextStyle(8.1, C["muted"], min_size=7))
    c.arrow(3.90, 2.55, 0.39, 0.22, C["grey"])
    c.arrow(8.95, 2.55, 0.39, 0.22, C["grey"])
    c.arrow(3.90, 4.90, 0.39, 0.22, C["grey"])
    c.arrow(8.95, 4.90, 0.39, 0.22, C["grey"])
    c.footer(total, f"Inventario actual: {m['plugins']} plugins · {m['skills']} skills · {m['commands']} comandos")
    slides.append(c)

    # 7 — dual engine
    c = Canvas(prs, 7, "Ejecución", "El proyecto puede elegir quién lleva las manos en el teclado", "AISDD y AIAD comparten la HU como unidad estable; cambia el motor de ejecución.")
    c.rect(0.66, 2.15, 5.72, 3.75, C["blue2"], C["blue2"], radius=0.16)
    c.pill(0.95, 2.47, 1.35, 0.38, "AI-ENGINE", C["blue"], C["white"], 8)
    c.text(0.95, 3.05, 4.95, 0.56, "AISDD", TextStyle(26, C["ink"], True))
    c.text(0.95, 3.68, 4.95, 0.65, "La IA implementa contra una especificación aprobada; la persona dirige, valida y decide.", TextStyle(12, C["ink2"], min_size=10))
    for i, t in enumerate(["open change", "implement", "close"]):
        c.pill(0.95 + i * 1.58, 4.72, 1.38, 0.42, t, C["white"], C["blue"], 7.2)
    c.text(0.95, 5.35, 4.9, 0.28, "Optimiza velocidad, consistencia y trazabilidad.", TextStyle(9.2, C["blue"], True))
    c.rect(6.95, 2.15, 5.72, 3.75, C["teal2"], C["teal2"], radius=0.16)
    c.pill(7.24, 2.47, 1.62, 0.38, "HUMAN-ENGINE", C["teal"], C["white"], 8)
    c.text(7.24, 3.05, 4.95, 0.56, "AIAD", TextStyle(26, C["ink"], True))
    c.text(7.24, 3.68, 4.95, 0.65, "La persona escribe; la IA explica, pregunta, prueba, revisa y acompaña cuando se la invoca.", TextStyle(12, C["ink2"], min_size=10))
    for i, t in enumerate(["think", "build", "review"]):
        c.pill(7.24 + i * 1.58, 4.72, 1.38, 0.42, t, C["white"], C["teal"], 7.2)
    c.text(7.24, 5.35, 4.9, 0.28, "Optimiza autoría, aprendizaje y oficio.", TextStyle(9.2, C["teal"], True))
    c.pill(5.42, 1.91, 2.48, 0.36, "BRIDGE · HU ↔ CHANGE", C["navy"], C["white"], 7.8)
    c.rect(0.66, 6.20, 12.01, 0.55, C["white"], C["line"], radius=0.08)
    c.text(0.96, 6.31, 11.42, 0.28, "Decisión por historia: no todo trabajo maximiza el mismo valor, y el motor puede cambiar a mitad sin perder la trazabilidad.", TextStyle(10.4, C["ink"], True, align="center"))
    c.footer(total)
    slides.append(c)

    # 8 — value chain
    c = Canvas(prs, 8, "Valor", "La evolución cierra los dos extremos que el código no cubre", "Antes del change está el cliente. Después del merge está la decisión de negocio.")
    nodes = [
        ("01", "Brief", "Qué necesita", C["blue"]),
        ("02", "Contrato", "RF · HU · DF", C["blue"]),
        ("03", "Diseño", "Arquitectura · Figma", C["cyan"]),
        ("04", "Entrega", "Change · código · tests", C["teal"]),
        ("05", "Evidencia", "Plan · estado · auditoría", C["amber"]),
        ("06", "Decisión", "KPI medido", C["red"]),
    ]
    for i, (n, t, b, color) in enumerate(nodes):
        xx = 0.55 + i * 2.10
        c.circle(xx + 0.58, 2.28, 0.62, color)
        c.text(xx + 0.58, 2.30, 0.62, 0.54, n, TextStyle(11, C["white"], True, align="center", valign="middle"))
        if i < len(nodes) - 1:
            c.arrow(xx + 1.35, 2.49, 0.60, 0.18, C["grey"])
        c.text(xx, 3.14, 1.78, 0.34, t, TextStyle(13, C["ink"], True, align="center"))
        c.text(xx, 3.62, 1.78, 0.62, b, TextStyle(8.8, C["ink2"], align="center", min_size=7.5))
    c.rect(0.65, 4.72, 5.85, 1.25, C["blue2"], C["blue2"], radius=0.12)
    c.text(0.94, 4.96, 1.70, 0.28, "EXTREMO CLIENTE", TextStyle(8.5, C["blue"], True))
    c.text(2.55, 4.88, 3.55, 0.52, "AIDD + AIBA convierten conversaciones en entregables revisables y firmables.", TextStyle(10.4, C["ink"], True, min_size=9, valign="middle"))
    c.rect(6.82, 4.72, 5.85, 1.25, C["amber2"], C["amber2"], radius=0.12)
    c.text(7.11, 4.96, 1.80, 0.28, "EXTREMO EVIDENCIA", TextStyle(8.5, C["amber"], True))
    c.text(8.90, 4.88, 3.36, 0.52, "AIBA explota auditoría y Jira sin disfrazar estimaciones de resultados medidos.", TextStyle(10.4, C["ink"], True, min_size=9, valign="middle"))
    c.footer(total)
    slides.append(c)

    # 9 — guarantees
    c = Canvas(prs, 9, "Escala y gobierno", "La metodología deja de depender de la memoria del equipo", "Las convenciones se convierten en decisiones explícitas, verificaciones y evidencia.")
    items = [
        ("DUDAS", "Pre-flight", "Bloqueantes siempre; preferencias configurables; decisiones persistidas.", C["blue"]),
        ("PARALELISMO", "3 modos", "atomic · waves · multilane, con garantías declaradas y barreras.", C["cyan"]),
        ("REPOS", "3 topologías", "mono · fraccionado · gobierno externo, sin inventar un repo padre.", C["teal"]),
        ("EVIDENCIA", "Auditoría", "Hashes, modelo, entradas, salidas y decisiones; sin guardar contenido sensible.", C["amber"]),
        ("CALIDAD", "12 checks CI", "Contratos, manifiestos, referencias, scripts, HTML y versiones.", C["red"]),
        ("DISTRIBUCIÓN", "6 plugins", "Semver independiente e instalación reversible por necesidad.", C["ink2"]),
    ]
    for i, (ey, title, body, accent) in enumerate(items):
        col, row = i % 3, i // 3
        x, y = 0.62 + col * 4.08, 2.13 + row * 2.08
        card(c, x, y, 3.82, 1.72, ey, title, body, accent)
    c.rect(0.62, 6.36, 12.06, 0.42, C["navy"], C["navy"], radius=0.08)
    c.text(0.86, 6.44, 11.58, 0.22, "Resultado: el método puede crecer en personas, repositorios y plataformas sin perder el contrato que lo hace reconocible.", TextStyle(9.5, C["white"], True, align="center"))
    c.footer(total)
    slides.append(c)

    # 10 — adoption
    c = Canvas(prs, 10, "Adopción", "No hace falta adoptar 34 skills el primer día", "La unidad de adopción es el problema que se quiere resolver, no el marketplace completo.")
    ladder = [
        (0.70, 5.25, 3.25, 0.98, "1 · FLUJO MÍNIMO", "AIDD + AISDD", "Definir y construir punta a punta", C["blue"]),
        (3.80, 4.28, 3.25, 1.95, "2 · CLIENTE REAL", "+ AIBA", "Firmables, pruebas, planificación, estado y KPIs", C["cyan"]),
        (6.90, 3.31, 2.76, 2.92, "3 · DISEÑO REAL", "+ AIFG", "Figma llega a la HU sin cargar el árbol entero", C["teal"]),
        (9.51, 2.34, 2.82, 3.89, "4 · ELECCIÓN", "AISDD ↔ AIAD", "Elegir IA o humano como motor por historia", C["amber"]),
    ]
    for x, y, w, h, ey, title, body, accent in ladder:
        c.rect(x, y, w, h, accent, accent, radius=0.12)
        c.text(x + 0.24, y + 0.18, w - 0.48, 0.24, ey, TextStyle(8, C["white"], True))
        c.text(x + 0.24, y + 0.53, w - 0.48, 0.46, title, TextStyle(17, C["white"], True, min_size=13))
        if h > 1.3:
            c.text(x + 0.24, y + 1.13, w - 0.48, h - 1.35, body, TextStyle(10.2, C["white"], min_size=8.5))
    c.text(0.70, 2.12, 5.65, 0.58, "Empieza por el núcleo. Añade superficie solo cuando aparece el coste que esa superficie evita.", TextStyle(19, C["ink"], True, min_size=16))
    c.text(0.70, 3.08, 5.50, 0.82, "Para un proyecto pequeño, AIDD + AISDD ya completan el recorrido. AIBA compensa cuando hay cliente, gobierno y reporting; AIFG cuando el diseño es contrato; AIAD cuando importa la autoría humana.", TextStyle(10.8, C["ink2"], min_size=9))
    c.footer(total)
    slides.append(c)

    # 11 — executive recommendation
    c = Canvas(prs, 11, "Decisión", "Qué conviene conservar, consolidar y hacer después", "Una recomendación ejecutiva, con la contrapartida a la vista.")
    card(c, 0.62, 2.12, 3.82, 3.52, "Conservar", "El núcleo", "Especificación antes de código, documentos como fuente de verdad, aprobación humana, changes pequeños y auditoría. Es la identidad del método.", C["blue"])
    card(c, 4.76, 2.12, 3.82, 3.52, "Consolidar", "La plataforma", "Presentar los seis plugins como capas de una cadena y no como un catálogo. Medir adopción por problema resuelto, no por skills instalados.", C["teal"])
    card(c, 8.90, 2.12, 3.82, 3.52, "Resolver", "La deuda", "Actualizar la metodología troncal para absorber paralelismo, topologías, Jira y las capas AIBA/AIAD/AIFG. Hoy los SKILL.md van por delante.", C["amber"])
    c.rect(0.62, 5.92, 12.10, 0.80, C["navy"], C["navy"], radius=0.10)
    c.text(0.94, 6.08, 2.60, 0.30, "DECISIÓN RECOMENDADA", TextStyle(8.5, C["cyan"], True))
    c.text(3.24, 6.02, 9.05, 0.42, "Adoptar por capas y financiar la reconciliación metodológica antes de escalar el relato fuera del equipo.", TextStyle(13, C["white"], True, min_size=10.5, valign="middle"))
    c.footer(total)
    slides.append(c)

    # 12 — appendix divider
    c = Canvas(prs, 12, "", dark=True)
    c.text(0.70, 0.70, 2.6, 0.25, "ANEXO", TextStyle(10, C["cyan"], True))
    c.text(0.70, 1.55, 8.80, 0.95, "El detalle detrás del relato", TextStyle(34, C["white"], True, min_size=28))
    c.text(0.70, 2.72, 8.20, 0.75, "Procedimientos, inventario de skills, diez mejoras, continuidades, costes y fuentes de contraste.", TextStyle(15, "C6D6E3", min_size=12))
    topics = ["A · Procedimientos", "B · Plugins y skills", "C · Mejoras", "D · Continuidad", "E · Costes y fuentes"]
    for i, topic in enumerate(topics):
        c.pill(0.70 + (i % 3) * 3.25, 4.25 + (i // 3) * 0.72, 2.85, 0.45, topic, "24435D", C["white"], 8)
    c.footer(total, "Anexo técnico · las siguientes slides conservan el detalle para consulta")
    slides.append(c)

    # 13 — side by side
    c = Canvas(prs, 13, "Anexo · procedimiento", "Los dos procedimientos, uno frente a otro", "Mismo núcleo de especificación; distinta amplitud operativa.")
    headers = ["DIMENSIÓN", "NATIVE-AI-SPECS v1.6", "MARKETPLACE v1.46.1"]
    xs, widths = [0.62, 3.27, 7.82], [2.55, 4.45, 4.86]
    for i, h in enumerate(headers):
        c.rect(xs[i], 2.02, widths[i], 0.56, C["navy"] if i == 0 else C["blue"] if i == 1 else C["teal"], radius=0.05)
        c.text(xs[i] + 0.12, 2.14, widths[i] - 0.24, 0.24, h, TextStyle(8.3, C["white"], True, align="center"))
    rows = [
        ("Piezas", "3 skills", "34 skills · 6 plugins"),
        ("Comandos", "7", "41"),
        ("Cobertura", "fases 3–5", "fases 0–5 + entrega + medición"),
        ("Roles", "4", "5 + cliente/negocio"),
        ("Paralelismo", "oleadas", "atomic · waves · multilane"),
        ("Repositorios", "uno", "mono · fraccionado · externo"),
        ("Distribución", "copia manual + check remoto", "marketplace · semver · release"),
        ("Control", "auditoría", "auditoría + 12 checks CI"),
    ]
    for r, row in enumerate(rows):
        yy = 2.66 + r * 0.48
        fill = C["white"] if r % 2 == 0 else "EEF1F2"
        for i, value in enumerate(row):
            c.rect(xs[i], yy, widths[i], 0.44, fill, fill, radius=0)
            c.text(xs[i] + 0.14, yy + 0.07, widths[i] - 0.28, 0.26, value, TextStyle(8.4, C["ink"] if i == 0 else C["ink2"], i == 0, align="left" if i == 0 else "center", min_size=7.2))
    c.rect(0.62, 6.62, 12.06, 0.26, C["amber2"], C["amber2"], radius=0.04)
    c.text(0.82, 6.65, 11.66, 0.17, "Lectura: la ligereza baja, pero la cobertura, la capacidad de escala y la interfaz con negocio suben.", TextStyle(7.7, C["ink"], True, align="center"))
    c.footer(total)
    slides.append(c)

    # 14 — inventory 1
    c = Canvas(prs, 14, "Anexo · inventario", "Del cliente al código: AIDD, AISDD y AIBA", "Los tres plugins que forman la cadena de entrega principal.")
    inventory = [
        (0.62, "AIDD · 9", "DEFINIR Y DISEÑAR", ["client-requirements", "requirements", "user-stories", "user-story-details", "prototype-architecture", "prototype", "style-guide", "architecture-proposal", "architecture"], C["blue"]),
        (4.55, "AISDD · 2", "ESPECIFICAR Y EJECUTAR", ["aisdd-specs", "aisdd-amend", "9 comandos: init · roadmap · open", "implement · amend · close", "lane · prototype-ux · uml"], C["teal"]),
        (8.48, "AIBA · 7", "NEGOCIO, ENTREGA Y MEDICIÓN", ["functional-design", "test-plan", "hu-review-plan", "project-plan", "sprint-planning", "status-report", "metrics"], C["amber"]),
    ]
    for x, name, role, skills, accent in inventory:
        c.rect(x, 2.05, 3.65, 4.52, C["white"], C["line"], radius=0.12)
        c.rect(x, 2.05, 3.65, 0.74, accent, accent, radius=0.12)
        c.text(x + 0.22, 2.20, 3.21, 0.27, name, TextStyle(12, C["white"], True))
        c.text(x + 0.22, 2.92, 3.21, 0.22, role, TextStyle(8, accent, True))
        for i, skill in enumerate(skills):
            yy = 3.35 + i * 0.34
            c.circle(x + 0.23, yy + 0.06, 0.10, accent)
            c.text(x + 0.48, yy, 2.92, 0.22, skill, TextStyle(8.5, C["ink2"], min_size=7.2))
    c.footer(total, "Nombres abreviados: todos los skills llevan el prefijo de su plugin.")
    slides.append(c)

    # 15 — inventory 2
    c = Canvas(prs, 15, "Anexo · inventario", "Autoría, diseño real y piezas compartidas", "Capas opcionales que cambian el modo de trabajo sin romper el flujo principal.")
    c.rect(0.62, 2.02, 5.73, 4.63, C["white"], C["line"], radius=0.12)
    c.text(0.90, 2.26, 4.90, 0.34, "AIAD · 11 skills", TextStyle(15, C["cyan"], True))
    groups = [
        ("PENSAR", "design · explain · rubber-duck"),
        ("CONSTRUIR", "tdd · test"),
        ("MEJORAR", "review"),
        ("FLUIR", "pair · bridge · unblock · save"),
        ("REGISTRAR", "journal"),
    ]
    for i, (group, skills) in enumerate(groups):
        yy = 2.91 + i * 0.63
        c.pill(0.90, yy, 1.15, 0.34, group, C["blue2"], C["blue"], 7.2)
        c.text(2.23, yy + 0.03, 3.66, 0.27, skills, TextStyle(9.2, C["ink2"], min_size=8))
    c.text(0.90, 6.11, 4.95, 0.25, "Principio: pull, no push. El humano conserva la autoría.", TextStyle(8.5, C["cyan"], True))
    card(c, 6.67, 2.02, 2.82, 2.02, "AIFG · 2", "Diseño → HU", "capture normaliza Figma por componente e historia; update re-captura y mide el impacto.", C["red"])
    card(c, 9.78, 2.02, 2.90, 2.02, "BOOSTERS · 3", "Visualizar", "UX crea prototipos; UML explica changes; Docs convierte Markdown en vistas HTML.", C["ink2"])
    c.rect(6.67, 4.40, 6.01, 2.25, C["navy"], C["navy"], radius=0.14)
    c.text(6.98, 4.70, 5.40, 0.28, "POR QUÉ SON ADITIVOS", TextStyle(9, C["cyan"], True))
    c.text(6.98, 5.18, 5.30, 0.78, "Sin AIFG, AISDD sigue leyendo la guía de estilos. Sin AIAD, AISDD sigue siendo el motor. Sin boosters, el contrato documental permanece: se pierde la vista, no la verdad.", TextStyle(11, C["white"], min_size=9))
    c.footer(total)
    slides.append(c)

    # 16 — improvements
    c = Canvas(prs, 16, "Anexo · mejoras", "Diez mejoras que cambian el proceso, no solo el tooling", "Agrupadas por el problema operativo que resuelven.")
    improvements = [
        ("01", "Definición ejecutable", "Brief, requisitos, HU y arquitectura como comandos", C["blue"]),
        ("02", "Interfaz con negocio", "Firmables, pruebas, plan, estado y KPI", C["amber"]),
        ("03", "Diseño hasta la HU", "Figma normalizado, cargado solo cuando aplica", C["red"]),
        ("04", "Paralelismo explícito", "Tres modos y garantías conocidas", C["cyan"]),
        ("05", "Multirepositorio", "Tres topologías y migración guiada", C["teal"]),
        ("06", "Retrabajo visible", "amend ejecuta y registra solo el delta", C["blue"]),
        ("07", "Métrica defendible", "Auditoría + worklog real de Jira", C["amber"]),
        ("08", "Autoverificación", "Doce checks evitan degradación silenciosa", C["red"]),
        ("09", "Segundo motor", "Human-first elegible por historia", C["cyan"]),
        ("10", "Distribución", "Plugins, semver y release reproducible", C["teal"]),
    ]
    for i, (n, title, body, accent) in enumerate(improvements):
        col, row = i % 2, i // 2
        x, y = 0.62 + col * 6.08, 2.02 + row * 0.91
        c.rect(x, y, 5.76, 0.72, C["white"], C["line"], radius=0.08)
        c.circle(x + 0.18, y + 0.14, 0.42, accent)
        c.text(x + 0.18, y + 0.15, 0.42, 0.36, n, TextStyle(8, C["white"], True, align="center", valign="middle"))
        c.text(x + 0.78, y + 0.11, 1.76, 0.24, title, TextStyle(9.4, C["ink"], True))
        c.text(x + 2.54, y + 0.11, 2.94, 0.42, body, TextStyle(7.7, C["ink2"], min_size=6.8, valign="middle"))
    c.footer(total)
    slides.append(c)

    # 17 — continuity
    c = Canvas(prs, 17, "Anexo · continuidad", "Lo que no se tocó es tan importante como lo nuevo", "La compatibilidad conceptual permite evolucionar sin obligar al equipo a reaprender el núcleo.")
    continuity = [
        "OpenSpec sigue debajo",
        "3 macrofases y 4 roles originales",
        "Pre-flight humano antes de actuar",
        "Auditoría append-only en openspec/audit",
        "AGENTS.md como ancla idempotente",
        "Referencias con carga diferida",
        "Alias native-ai siguen funcionando",
        "9 principios fundacionales conservados",
    ]
    for i, item in enumerate(continuity):
        col, row = i % 2, i // 2
        x, y = 0.72 + col * 6.10, 2.10 + row * 0.88
        c.circle(x, y + 0.06, 0.35, C["teal"])
        c.text(x, y + 0.04, 0.35, 0.29, "OK", TextStyle(6.5, C["white"], True, align="center", valign="middle"))
        c.text(x + 0.55, y, 5.18, 0.46, item, TextStyle(10.5, C["ink"], True, min_size=9, valign="middle"))
    c.rect(0.72, 5.85, 11.88, 0.78, C["teal2"], C["teal2"], radius=0.10)
    c.text(0.99, 6.01, 11.34, 0.38, "Prueba de continuidad: ninguna mejora necesita negar «especificar antes de implementar»; amplía dónde empieza, dónde termina o quién ejecuta.", TextStyle(10.7, C["ink"], True, align="center", valign="middle"))
    c.footer(total)
    slides.append(c)

    # 18 — honesty and sources
    c = Canvas(prs, 18, "Anexo · honestidad", "El precio de la evolución — y la evidencia usada", "Más capacidad implica más superficie, deuda documental y una adopción que debe ser deliberada.")
    card(c, 0.62, 2.05, 3.72, 2.18, "Coste 1", "Más piezas", "3 skills y 7 comandos se leen en una tarde. 34 skills y 6 plugins exigen arquitectura de adopción, ownership y formación.", C["amber"])
    card(c, 4.52, 2.05, 3.72, 2.18, "Coste 2", "Dos velocidades", "La metodología troncal va por detrás de capacidades ya operativas en AISDD, AIBA, AIAD y AIFG.", C["red"])
    card(c, 8.42, 2.05, 4.26, 2.18, "Respuesta", "Adopción por capas", "Núcleo primero; negocio, diseño y human-first solo cuando el proyecto puede convertirlos en valor.", C["teal"])
    c.rect(0.62, 4.55, 12.06, 1.70, C["white"], C["line"], radius=0.12)
    c.text(0.90, 4.80, 2.05, 0.25, "FUENTES DE CONTRASTE", TextStyle(8.5, C["blue"], True))
    sources = "ai-native.md v2.0 · native-ai-specs v1.6.0 · marketplace.json · plugin.json · 34 SKILL.md · metodologías AIDD/AISDD/AIBA/AIAD · 12 check_*.py"
    c.text(2.72, 4.72, 9.53, 0.56, sources, TextStyle(10.1, C["ink"], True, min_size=8.5, valign="middle"))
    c.text(0.90, 5.48, 11.38, 0.42, "Corte de datos: 09/09/2026 · Marketplace VERSION 1.46.1 · Las cifras son conteos; las valoraciones cualitativas son interpretación del equipo.", TextStyle(8.8, C["muted"], min_size=7.5, align="center"))
    c.rect(0.62, 6.47, 12.06, 0.38, C["navy"], C["navy"], radius=0.06)
    c.text(0.85, 6.54, 11.60, 0.20, "CONCLUSIÓN  ·  Native AI conserva su promesa original y ahora puede sostenerla desde la conversación con el cliente hasta la decisión basada en evidencia.", TextStyle(8.7, C["white"], True, align="center"))
    c.footer(total)
    slides.append(c)

    return slides


def add_document_properties(prs: Presentation):
    props = prs.core_properties
    props.title = "Evolución de Native AI — resumen ejecutivo"
    props.subject = "Comparativa entre la metodología original, native-ai-specs v1.6.0 y el marketplace"
    props.author = "NTT DATA Spain GDN-e"
    props.keywords = "Native AI, AIDD, AISDD, AIBA, AIAD, AIFG, ejecutivo"
    props.comments = "Generado de forma reproducible con generate.py. Issue #47."


def save_outputs(prs: Presentation, slides: Iterable[Canvas]):
    add_document_properties(prs)
    prs.save(PPTX_PATH)
    RENDERED.mkdir(exist_ok=True)
    images = []
    for canvas in slides:
        path = RENDERED / f"slide-{canvas.number:02d}.png"
        canvas.image.save(path, quality=95)
        images.append(canvas.image.convert("RGB"))
    # 1600 px / 120 dpi = 13,333 in: mismo tamaño físico que el lienzo 16:9 del PPTX.
    images[0].save(PDF_PATH, "PDF", resolution=120.0, save_all=True, append_images=images[1:])
    thumb_w, thumb_h = 384, 216
    cols, rows = 3, math.ceil(len(images) / 3)
    sheet = Image.new("RGB", (cols * thumb_w + (cols + 1) * 18, rows * thumb_h + (rows + 1) * 18), rgba(C["soft"])[0:3])
    for i, source in enumerate(images):
        thumb = source.copy()
        thumb.thumbnail((thumb_w, thumb_h), Image.Resampling.LANCZOS)
        x = 18 + (i % cols) * (thumb_w + 18)
        y = 18 + (i // cols) * (thumb_h + 18)
        sheet.paste(thumb, (x, y))
    sheet.save(PREVIEW_PATH, quality=92)


def verify(prs: Presentation, slides: list[Canvas]):
    if len(slides) != 18 or len(prs.slides) != 18:
        raise RuntimeError("El deck debe contener exactamente 18 diapositivas (11 ejecutivas + 7 de anexo).")
    for slide_no, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width or shape.top + shape.height > prs.slide_height:
                raise RuntimeError(f"Elemento fuera del lienzo en slide {slide_no}: {shape.name}")
    required = [PPTX_PATH, PDF_PATH, PREVIEW_PATH]
    missing = [str(path) for path in required if not path.exists() or path.stat().st_size == 0]
    if missing:
        raise RuntimeError(f"Artefactos no generados: {missing}")


def main():
    m = metrics()
    if RENDERED.exists():
        shutil.rmtree(RENDERED)
    prs = Presentation()
    prs.slide_width = Inches(W_IN)
    prs.slide_height = Inches(H_IN)
    slides = build_slides(prs, m)
    save_outputs(prs, slides)
    verify(prs, slides)
    print(f"Generados {len(slides)} slides: {PPTX_PATH.name}, {PDF_PATH.name} y {PREVIEW_PATH.name}")
    print(f"Métricas: {m['plugins']} plugins · {m['skills']} skills · {m['commands']} comandos · {m['checks']} checks")


if __name__ == "__main__":
    main()
